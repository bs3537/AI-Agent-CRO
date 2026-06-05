"""Best-effort extraction for chatbot uploads."""
from __future__ import annotations

import csv
import io
import json
import logging
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger("sma_monitor.chat.files")

MAX_UPLOAD_BYTES = 15 * 1024 * 1024
MAX_EXTRACTED_CHARS = 12_000
TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl",
    ".html", ".htm", ".xml", ".log", ".py", ".js", ".ts", ".tsx",
}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}


@dataclass(frozen=True)
class ChatAttachment:
    filename: str
    content_type: str
    byte_size: int
    n_chars: int
    text: str
    parser: str


class ChatFileError(RuntimeError):
    pass


def parse_upload(filename: str, content_type: str | None, content: bytes) -> ChatAttachment:
    if len(content) > MAX_UPLOAD_BYTES:
        size_mb = MAX_UPLOAD_BYTES // (1024 * 1024)
        raise ChatFileError(f"{filename} exceeds the {size_mb} MB chat limit")
    safe = Path(filename or "upload").name
    suffix = Path(safe).suffix.lower()
    ctype = content_type or "application/octet-stream"

    if suffix in TEXT_SUFFIXES:
        text = _text_from_bytes(content, suffix)
        return _attachment(safe, ctype, content, text, "local_text")
    if suffix in PDF_SUFFIXES:
        text = _extract_pdf(content)
        return _attachment(safe, ctype, content, text, "local_pdf")
    if suffix in DOCX_SUFFIXES:
        return _attachment(safe, ctype, content, _extract_docx(content), "local_docx")
    if suffix in EXCEL_SUFFIXES:
        return _attachment(safe, ctype, content, _extract_excel(content, suffix), "local_excel")
    if ctype.startswith("image/"):
        raise ChatFileError("image chat uploads are not supported")
    raise ChatFileError(f"unsupported chat upload type: {suffix or ctype}")


def _attachment(
    filename: str,
    content_type: str,
    content: bytes,
    text: str,
    parser: str,
) -> ChatAttachment:
    text = _truncate(text.strip())
    return ChatAttachment(
        filename=filename,
        content_type=content_type,
        byte_size=len(content),
        n_chars=len(text),
        text=text,
        parser=parser,
    )


def _text_from_bytes(content: bytes, suffix: str) -> str:
    text = content.decode("utf-8", errors="replace")
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        rows = csv.reader(io.StringIO(text), delimiter=delimiter)
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in list(rows)[:300])
    if suffix in {".json", ".jsonl"}:
        try:
            return json.dumps(json.loads(text), indent=2)[:MAX_EXTRACTED_CHARS]
        except json.JSONDecodeError:
            return text
    return text


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ChatFileError("PDF extraction needs pypdf") from e
    try:
        reader = PdfReader(io.BytesIO(content))
        return "\n".join((page.extract_text() or "") for page in reader.pages[:80]).strip()
    except Exception as e:  # noqa: BLE001
        log.warning("chat_pdf_extract_failed", extra={"err": str(e)[:200]})
        return ""


def _extract_docx(content: bytes) -> str:
    try:
        import docx
    except ImportError as e:
        raise ChatFileError("DOCX extraction needs python-docx") from e
    document = docx.Document(io.BytesIO(content))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_excel(content: bytes, suffix: str) -> str:
    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as e:
            raise ChatFileError("Excel extraction needs openpyxl") from e
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        parts: list[str] = []
        for ws in wb.worksheets[:8]:
            parts.append(f"### Sheet: {ws.title}")
            for row in ws.iter_rows(max_row=120, max_col=30, values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    try:
        import xlrd
    except ImportError as e:
        raise ChatFileError("XLS extraction needs xlrd") from e
    book = xlrd.open_workbook(file_contents=content)
    parts = []
    for sheet in book.sheets()[:8]:
        parts.append(f"### Sheet: {sheet.name}")
        for r in range(min(sheet.nrows, 120)):
            row = [str(sheet.cell_value(r, c)) for c in range(min(sheet.ncols, 30))]
            if any(cell.strip() for cell in row):
                parts.append(" | ".join(row))
    return "\n".join(parts)


def _truncate(text: str) -> str:
    return text[:MAX_EXTRACTED_CHARS] if len(text) > MAX_EXTRACTED_CHARS else text
